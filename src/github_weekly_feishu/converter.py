from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
LINK_RE = re.compile(r"(?<!!)\[([^\]]+)\]\(([^)]+)\)")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")
ORDERED_LIST_RE = re.compile(r"^\d+\.\s+(.+)$")


def clean_inline(text: str) -> str:
    """Convert lightweight Markdown inline syntax to plain document text."""
    text = IMAGE_RE.sub("", text)
    text = LINK_RE.sub(r"\1（\2）", text)
    for marker in ("**", "__", "`"):
        text = text.replace(marker, "")
    return text.strip()


def _is_table_separator(cells: Iterable[str]) -> bool:
    return all(cell and set(cell) <= {"-", ":", " "} for cell in cells)


def _add_markdown_table(doc: Document, rows: list[str]) -> None:
    parsed: list[list[str]] = []
    for row in rows:
        cells = [clean_inline(cell.strip()) for cell in row.strip("|").split("|")]
        if _is_table_separator(cells):
            continue
        parsed.append(cells)

    if not parsed:
        return

    width = max(len(row) for row in parsed)
    table = doc.add_table(rows=len(parsed), cols=width)
    table.style = "Table Grid"

    for row_index, cells in enumerate(parsed):
        for cell_index in range(width):
            value = cells[cell_index] if cell_index < len(cells) else ""
            table.cell(row_index, cell_index).text = value
            if row_index == 0:
                for paragraph in table.cell(row_index, cell_index).paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def _add_image(doc: Document, assets_root: Path, alt: str, rel_path: str) -> None:
    image_path = assets_root / rel_path
    if not image_path.exists() or not image_path.is_file():
        return

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.2))

    if alt:
        caption = doc.add_paragraph(clean_inline(alt))
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in caption.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(100, 100, 100)


def build_doc(source: Path, output: Path, assets_root: Path | None = None) -> Path:
    """Build a Feishu/Word friendly DOCX document from a Markdown source file."""
    source = Path(source)
    output = Path(output)
    root = Path(assets_root) if assets_root is not None else source.parent

    if not source.exists():
        raise FileNotFoundError(f"Markdown source does not exist: {source}")
    if not source.is_file():
        raise ValueError(f"Markdown source is not a file: {source}")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    table_buffer: list[str] = []

    def flush_table() -> None:
        nonlocal table_buffer
        if table_buffer:
            _add_markdown_table(doc, table_buffer)
            table_buffer = []

    for raw in source.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if not line:
            flush_table()
            continue

        if line.startswith("|"):
            table_buffer.append(line)
            continue

        flush_table()

        image = IMAGE_RE.fullmatch(line.strip())
        if image:
            _add_image(doc, root, image.group(1), image.group(2))
            continue

        heading = HEADING_RE.match(line)
        if heading:
            level = min(len(heading.group(1)) - 1, 4)
            doc.add_heading(clean_inline(heading.group(2)), level=level)
            continue

        if line == "---":
            doc.add_paragraph("")
            continue

        if line.startswith("> "):
            paragraph = doc.add_paragraph(clean_inline(line[2:]))
            paragraph.style = "Intense Quote"
            continue

        if line.startswith(("- ", "* ", "+ ")):
            doc.add_paragraph(clean_inline(line[2:]), style="List Bullet")
            continue

        ordered_item = ORDERED_LIST_RE.match(line)
        if ordered_item:
            doc.add_paragraph(clean_inline(ordered_item.group(1)), style="List Number")
            continue

        doc.add_paragraph(clean_inline(line))

    flush_table()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output
