from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from docx import Document

from github_weekly_feishu import build_doc
from github_weekly_feishu.cli import main


class ConverterTests(unittest.TestCase):
    def test_build_doc_converts_common_weekly_markdown(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "weekly.md"
            output = root / "weekly.docx"
            source.write_text(
                "\n".join(
                    [
                        "# Weekly",
                        "",
                        "> Quote with **bold** text",
                        "",
                        "| Project | URL |",
                        "| --- | --- |",
                        "| Demo | [repo](https://github.com/example/demo) |",
                        "",
                        "- First item",
                        "1. Numbered item",
                        "![missing](assets/missing.png)",
                    ]
                ),
                encoding="utf-8",
            )

            generated = build_doc(source, output)

            self.assertEqual(generated, output)
            self.assertTrue(output.exists())

            doc = Document(output)
            texts = [paragraph.text for paragraph in doc.paragraphs]
            self.assertIn("Weekly", texts)
            self.assertIn("Quote with bold text", texts)
            self.assertIn("First item", texts)
            self.assertIn("Numbered item", texts)
            self.assertEqual(len(doc.tables), 1)
            self.assertEqual(doc.tables[0].cell(1, 1).text, "repo（https://github.com/example/demo）")

    def test_cli_reports_missing_input(self) -> None:
        code = main(["missing.md"])
        self.assertEqual(code, 2)


if __name__ == "__main__":
    unittest.main()
